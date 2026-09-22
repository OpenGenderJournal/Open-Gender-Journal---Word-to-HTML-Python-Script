pip install mammoth beautifulsoup4 python-docx
import mammoth
from bs4 import BeautifulSoup
from docx import Document

INPUT_FILE = r"C:\Users\pfist\Desktop\OGJ_Template_Rezension_Deutsch.docx"
OUTPUT_FILE = r"C:\Users\pfist\Desktop\output.html"
TEMP_FILE = r"C:\Users\pfist\Desktop\temp_cleaned.docx"


style_map = """
p[style-name='OGJ: Überschrift 1'] => h1:fresh
p[style-name='OGJ: Überschrift 2'] => h1:fresh
p[style-name='OGJ: Überschrift 3'] => h1:fresh
p[style-name='OGJ: Referenzen'] => h1:fresh
p[style-name='OGJ: Manuskript Titel'] => h1:fresh
b => strong
"""


def remove_headers_footers(doc):

    for section in doc.sections:

        header = section.header
        for paragraph in list(header.paragraphs):

            if paragraph.style.name.startswith("OGJ: Kopfzeile"):

                p = paragraph._element
                p.getparent().remove(p)

        footer = section.footer
        for paragraph in list(footer.paragraphs):

            if paragraph.style.name.startswith("OGJ: Fußzeile"):

                p = paragraph._element
                p.getparent().remove(p)


def add_reference_ids(html):

    soup = BeautifulSoup(html, "html.parser")

    reference_section = None

    for h in soup.find_all("h1"):

        if "referenzen" in h.text.lower() or "literatur" in h.text.lower():

            reference_section = h
            break

    if reference_section is None:
        return str(soup)

    counter = 1
    current = reference_section.find_next_sibling()

    while current:

        if current.name == "h1":
            break

        if current.name == "p":

            ref_id = f"CIT{counter:04d}"

            anchor = soup.new_tag("a")
            anchor["id"] = ref_id

            current.insert_before(anchor)

            counter += 1

        current = current.find_next_sibling()

    return str(soup)


def convert_docx():

    doc = Document(INPUT_FILE)

    remove_headers_footers(doc)

    doc.save(TEMP_FILE)

    with open(TEMP_FILE, "rb") as docx_file:

        result = mammoth.convert_to_html(
            docx_file,
            style_map=style_map
        )

        html = result.value

    html = add_reference_ids(html)

    full_html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<title>OGJ Export</title>

<link rel="stylesheet" type="text/css" href="ogjscreen.css">

<style>

/* Tabellenlinien erzwingen */

#paper table {{
border-collapse: collapse;
width: 100%;
margin: 20px 0;
}}

#paper td, #paper th {{
border: 1px solid black;
padding: 6px;
}}

#paper th {{
background-color: #eeeeee;
}}

</style>

</head>

<body id="paper">

{html}

</body>
</html>
"""

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:

        f.write(full_html)

    print("✅ Konvertierung abgeschlossen:")
    print(OUTPUT_FILE)


if __name__ == "__main__":
    convert_docx()
